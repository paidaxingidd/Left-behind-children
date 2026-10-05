"""Assemble all results into one Excel workbook (one sheet per analysis)."""
import json, os, pandas as pd, numpy as np
from lbc_core import PREFIX
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook(); wb.remove(wb.active)
def sheet(name, df, note=None):
    ws = wb.create_sheet(name)
    if note:
        ws.append([note]); ws['A1'].font = Font(italic=True); ws.append([])
    ws.append(list(df.columns))
    for c in ws[ws.max_row]: c.font = Font(bold=True)
    for row in df.itertuples(index=False):
        ws.append([float(v) if isinstance(v, (np.floating, float)) else (int(v) if isinstance(v, (np.integer,)) else v) for v in row])
    for i, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(60, int(df[col].astype(str).str.len().max()) + 2))
    return ws

A = pd.read_csv(PREFIX + 'res_A_sensitivity.csv')
sheet('A_D_E_sensitivity', A.round(3), '主模型（家庭功能总分→生活满意度→学习投入；协变量性别、监护、离异）在不同样本/量表/协变量下的标准化系数与 5000 次 bootstrap 百分位 CI。正式版：b13 按稿件口径（311 人取 2024 年文件值，111 人插补；FAD 30 题）。')

cfa = json.load(open(PREFIX + 'res_A3_fad_cfa.json')); w = json.load(open(PREFIX + 'res_A3_wording_models.json')); ri = json.load(open(PREFIX + 'res_A3_riifa.json'))
rows = [dict(model='One factor', **w['one_factor']), dict(model='Two wording factors (neg / pos), r = %.2f' % w['r_neg_pos'], **w['two_wording_factors']),
        dict(model='Five substantive factors (manuscript model)', **cfa['none']['fit']),
        dict(model='Five factors + free-loading method factor on 19 negatively worded items (orthogonal)', **cfa['neg']['fit']),
        dict(model='Five factors + random-intercept (acquiescence) factor, loadings fixed ±1', **{k: round(float(np.mean([ri[j]['fit'][k] for j in ri])), 3) for k in ['chi2', 'df', 'cfi', 'tli', 'rmsea', 'srmr']})]
sheet('A3_FAD_CFA', pd.DataFrame(rows), '第一个插补集，最大似然，FAD 30 题。措辞方法因子模型：拟合略升（CFI .85→.88），但方法因子平均解释消极措辞题目 19% 方差，家庭规则因子实质载荷消失（.38/−.31/.11/.08/−.11）；再加正向措辞因子后问题解决因子载荷全部变号（不可解释）；随机截距模型给出不可容许解（因子相关 >1）。原因：每个分量表的题目措辞方向完全一致，措辞效应与分量表结构在统计上不可分。')
ld = pd.DataFrame({'factor': list(cfa['none']['loadings']), 'five_factor_loadings': [str(v) for v in cfa['none']['loadings'].values()], 'with_method_factor': [str(v) for v in cfa['neg']['loadings'].values()]})
sheet('A3_FAD_loadings', ld)
fc = pd.DataFrame({'pair': list(cfa['none']['factor_corr']), 'five_factor': list(cfa['none']['factor_corr'].values()), 'with_method_factor': list(cfa['neg']['factor_corr'].values())})
sheet('A3_FAD_factor_corr', fc)

B = pd.read_csv(PREFIX + 'res_B_h3.csv')
lab = {'a_': 'a path (unique) – ', 'cp_': "c' (unique) – ", 'ind_': 'indirect (joint model) – ', 'sep_': 'indirect (separate model, Table 3) – ',
       'sepdiff_': 'difference of separate-model indirects: ', 'jointdiff_': 'difference of joint-model indirects: '}
def describe(p):
    for k, v in lab.items():
        if p.startswith(k): return v + p[len(k):]
    return {'b': 'b path (joint five-dimension model)', 'b2': 'b path (two-composite model)', 'diff_aff_org_joint': 'indirect(affective) − indirect(organisational), two-composite model',
            'diff_aff_org_sep': 'mean indirect(affective dims) − mean indirect(organisational dims), separate models'}.get(p, p)
B.insert(1, 'description', B.param.map(describe))
sheet('B_H3_joint_models', B.round(3), '假设 3：五个维度同时进入一个中介模型（a、c′ 为控制其他维度后的独特效应）；情感/组织两个合成变量（各维度 z 分均值）；差异的 90% CI 即 TOST 等价检验的判定区间（区间落在 ±δ 内即可宣称等价）。')

C = pd.read_csv(PREFIX + 'res_C_lbc_compare.csv')
keep = ['label', 'n', 'n_lbc'] + [c for c in C.columns if c.startswith(('a_diff', 'b_LBC', 'b_non', 'b_diff', 'cp_diff', 'ind_LBC', 'ind_non', 'ind_diff'))]
sheet('C_LBC_vs_nonLBC', C[keep].round(3), '留守身份调节 a、b、c′ 三条路径（535 人；协变量性别、年龄、离异；分层 bootstrap 5000 次）。C1 加年龄×生活满意度、年龄×家庭功能、年龄×留守交互；C2 年级固定效应；C3 六年级子样本；C4 留守组按年级加权到非留守组的年级分布；C5 两组分别标准化后分别建模。')

F1 = pd.read_csv(PREFIX + 'res_F1_disattenuated.csv'); sheet('F1_reliability_corrected', F1.round(3), '单指标信度校正：用各插补集的 Cronbach α 对三个总分的相关做失真校正后再估计路径（协变量视为无误差）；bootstrap 5000 次。')
try:
    F2 = pd.read_csv(PREFIX + 'res_F2_latent_sem.csv'); fit = json.load(open(PREFIX + 'res_F2_fit.json'))
    sheet('F2_latent_SEM', F2.round(3), '完整潜变量 SEM：FF 为五个 FAD 一阶因子之上的二阶因子（30 题），LS 潜变量（5 题），LE 潜变量（活力/奉献/专注三个小包）；协变量预测 LS 与 LE。拟合（20 插补集均值）：CFI %.3f, TLI %.3f, RMSEA %.3f, SRMR %.3f；bootstrap %d 次。注意：二阶因子由消极措辞的三个分量表主导（载荷 EM .75, ES .98, FR .84 vs PC .34, PS .41）。' % (fit['fit_pooled']['cfi'], fit['fit_pooled']['tli'], fit['fit_pooled']['rmsea'], fit['fit_pooled']['srmr'], fit['n_boot']))
except FileNotFoundError:
    pass
T3 = pd.read_csv(PREFIX + 'table3.csv'); sheet('Table3_interim_check', T3.round(3), '流程验证：按稿件口径重跑表 3，六行均应与稿件表 3 一致（小数点后 2 位）。')
wb.save(PREFIX + 'LBC_supplementary_analyses_2026-09-28.xlsx'); print('saved')
